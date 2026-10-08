"""Import/export en masse du portefeuille Configurateur via Excel - voir
docs/specs/portfolio_import.md. Demande de l'utilisateur (2026-09-28) :
l'app n'a pas de memoire entre sessions (`st.session_state` est reinitialise
a chaque rechargement Streamlit) - au lieu de ressaisir chaque projet a la
main dans le formulaire, l'utilisateur maintient sa propre liste de configs
dans un fichier Excel qu'il reupload a chaque session. Un seul onglet,
lecture **par recherche de libelle de colonne** (ligne d'en-tete), pas par
position fixe - meme convention que `bp_parser.py`
(docs/specs/bp_parsing.md), ici appliquee a un format qu'on controle
entierement (pas un export Aurora tiers) donc pas de recherche de la ligne
d'en-tete elle-meme : toujours la 1ere ligne du seul onglet lu."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import contract_overlay
from .bp_parser import _normalize, load_grid
from .contract_overlay import ContractStructure
from .portfolio import ProjectConfig

SHEET_NAME = "Projects"
LEGEND_SHEET_NAME = "Legend"

_CONTRACT_LABEL_TO_KIND = {
    "full merchant": contract_overlay.FULL_MERCHANT,
    "floor": contract_overlay.FLOOR,
    "tolling": contract_overlay.TOLLING,
}

# Valeurs litterales attendues (pas de mapping "libelle humain -> code" comme
# pour _CONTRACT_LABEL_TO_KIND : ce sont les memes 3 modes que le selectbox du
# Configurateur, voir ui/configurateur_tab.py). Une valeur hors de cet
# ensemble doit lever une erreur claire ici, pas planter loin plus tard dans
# dev_case.connection_capex_keur avec un ValueError generique illisible en
# prod (trouve le 2026-09-29 : un utilisateur avait tape "distance" au lieu
# de "distance_rte" dans son fichier, l'app a plante sur Streamlit Cloud avec
# une trace opaque - "zero zero silencieux" s'applique aussi aux erreurs, pas
# seulement aux valeurs manquantes).
_VALID_CONNECTION_CAPEX_MODES = {"library", "manual", "distance_rte"}
_VALID_REPOWERING_YEAR_MODES = {"auto", "manual", "soh"}

# (libelle de colonne, description, valeurs attendues) - sert a la fois de
# feuille "Legend" du template et de documentation inline ici. Ordre =
# ordre des colonnes generees.
_COLUMN_SPECS: list[tuple[str, str, str]] = [
    ("Project name", "Free text, must be unique in the file", "e.g. Site A"),
    ("BESS duration (h)", "Required", "2 or 4"),
    ("Voltage (Tension)", "Required", "HTA, HTB1, HTB2 or HTB3"),
    ("TURPE type", "Required", "Classique, Injection or Soutirage"),
    ("Gabarit", "Required", "Yes or No"),
    ("COD year", "Required", "e.g. 2028"),
    ("Power (MW)", "Required", "e.g. 10"),
    ("Operating life (years)", "Optional, default 20", "e.g. 20"),
    (
        "Contract structure",
        "Optional, default Full merchant",
        "Full merchant, Floor or Tolling",
    ),
    ("Contract price (k€/MW/yr)", "Used if Floor or Tolling", "e.g. 80"),
    ("Contract duration (years)", "Used if Floor or Tolling", "e.g. 10"),
    (
        "Aggregator sharing above floor (%)",
        "Used if Floor, default 0",
        "e.g. 0",
    ),
    ("ORO", "Optional, default No", "Yes or No"),
    ("ORO curtailment hours", "Used if ORO=Yes, default 3000", "500-4000, step 500"),
    ("Repowering enabled", "Optional, default Yes", "Yes or No"),
    (
        "Repowering year mode",
        "Used if Repowering enabled=Yes, default manual",
        "auto, manual or soh (SoH-triggered, Aurora method)",
    ),
    (
        "Repowering year (manual)",
        "Used if Repowering year mode=manual, default 15",
        "op-year since COD, e.g. 15",
    ),
    (
        "Grid connection CAPEX mode",
        "Optional, default library",
        "library, manual or distance_rte",
    ),
    ("Manual grid connection CAPEX (k€)", "Used if mode=manual", "e.g. 2000"),
    ("Distance to RTE substation (km)", "Used if mode=distance_rte", "e.g. 5"),
    (
        "Land lease OPEX (k€/yr)",
        "Optional, default 0 (uses Aurora's generic estimate) - REPLACES it when set",
        "e.g. 0",
    ),
    (
        "Land lease indexation (%/yr)",
        "Optional, default 0 (flat rent) - indexes the land lease from op-year 2: "
        "rent year N = rent year 1 x (1 + rate)^(N-1)",
        "e.g. 2",
    ),
    ("Gearing override (%)", "Optional, blank = model default", "e.g. 70"),
    ("Interest rate override (%)", "Optional, blank = model default", "e.g. 5"),
    ("DSA override (k€)", "Optional, blank = model default", "e.g. 500"),
    ("DEVEX override (k€)", "Optional, blank = model default", "e.g. 300"),
    ("Carry months override", "Optional, blank = model default", "e.g. 18"),
    ("Carry rate override (%)", "Optional, blank = model default", "e.g. 8"),
    (
        "TURPE 50% reduction",
        "Optional, default No - HTB1/HTB2/HTB3 only, see docs/specs/turpe_50pct_reduction.md",
        "Yes or No",
    ),
]
_COLUMNS = [label for label, _, _ in _COLUMN_SPECS]
_REQUIRED_COLUMNS = [
    "Project name",
    "BESS duration (h)",
    "Voltage (Tension)",
    "TURPE type",
    "Gabarit",
    "COD year",
    "Power (MW)",
]


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _to_bool(value, *, default: bool) -> bool:
    if _is_blank(value):
        return default
    if isinstance(value, bool):
        return value
    return _normalize(value) in ("yes", "true", "1", "oui", "x")


def _to_int(value, *, default: int | None) -> int | None:
    if _is_blank(value):
        return default
    return int(value)


def _to_float(value, *, default: float | None) -> float | None:
    if _is_blank(value):
        return default
    return float(value)


def _to_pct(value, *, default: float | None) -> float | None:
    """Colonnes "... (%)" du template : saisies en points (70 = 70%), pas en
    fraction - conversion /100 ici, jamais demandee a l'utilisateur."""
    value = _to_float(value, default=None)
    return default if value is None else value / 100.0


def _header_index(header_row: list) -> dict[str, int]:
    index = {_normalize(cell): i for i, cell in enumerate(header_row) if not _is_blank(cell)}
    missing = [c for c in _REQUIRED_COLUMNS if _normalize(c) not in index]
    if missing:
        raise ValueError(
            f"Missing required column(s) in the '{SHEET_NAME}' sheet: {', '.join(missing)}. "
            "Download the template to see the expected format."
        )
    return index


def _cell(row: list, index: dict[str, int], label: str):
    col = index.get(_normalize(label))
    if col is None or col >= len(row):
        return None
    return row[col]


def parse_portfolio_excel(file_or_path) -> list[ProjectConfig]:
    """Lit l'onglet `Projects` d'un classeur uploade -> une `ProjectConfig`
    par ligne non vide. Colonnes requises (voir `_REQUIRED_COLUMNS`) leve un
    `ValueError` explicite si absentes ou vides sur une ligne - jamais de
    valeur inventee pour un champ obligatoire (brief section 7, "zero zero
    silencieux"). Les colonnes optionnelles absentes du fichier retombent sur
    le defaut du dataclass `ProjectConfig`, comme si la colonne existait mais
    etait vide partout."""
    grid = load_grid(file_or_path, sheet_name=SHEET_NAME)
    if not grid:
        raise ValueError(f"The '{SHEET_NAME}' sheet is empty.")
    index = _header_index(grid[0])

    configs: list[ProjectConfig] = []
    for row_number, row in enumerate(grid[1:], start=2):
        if all(_is_blank(v) for v in row):
            continue
        name = _cell(row, index, "Project name")
        try:
            for label in _REQUIRED_COLUMNS:
                if _is_blank(_cell(row, index, label)):
                    raise ValueError(f"'{label}' is required.")

            contract_label = _cell(row, index, "Contract structure")
            contract_kind = (
                contract_overlay.FULL_MERCHANT
                if _is_blank(contract_label)
                else _CONTRACT_LABEL_TO_KIND[_normalize(contract_label)]
            )
            oro_requested = _to_bool(_cell(row, index, "ORO"), default=False)
            repowering_enabled = _to_bool(_cell(row, index, "Repowering enabled"), default=True)
            repowering_year_mode = _cell(row, index, "Repowering year mode")
            repowering_year_mode = (
                "manual" if _is_blank(repowering_year_mode) else _normalize(repowering_year_mode)
            )
            if repowering_year_mode not in _VALID_REPOWERING_YEAR_MODES:
                raise ValueError(
                    f"'Repowering year mode' must be one of: "
                    f"{', '.join(sorted(_VALID_REPOWERING_YEAR_MODES))} - got '{repowering_year_mode}'."
                )
            connection_capex_mode = _cell(row, index, "Grid connection CAPEX mode")
            connection_capex_mode = (
                "library" if _is_blank(connection_capex_mode) else _normalize(connection_capex_mode)
            )
            if connection_capex_mode not in _VALID_CONNECTION_CAPEX_MODES:
                raise ValueError(
                    f"'Grid connection CAPEX mode' must be one of: "
                    f"{', '.join(sorted(_VALID_CONNECTION_CAPEX_MODES))} - got "
                    f"'{connection_capex_mode}'."
                )

            configs.append(
                ProjectConfig(
                    name=str(name).strip(),
                    duree_h=int(_cell(row, index, "BESS duration (h)")),
                    tension=str(_cell(row, index, "Voltage (Tension)")).strip(),
                    turpe_type=str(_cell(row, index, "TURPE type")).strip(),
                    gabarit=_to_bool(_cell(row, index, "Gabarit"), default=False),
                    cod_year=int(_cell(row, index, "COD year")),
                    power_mw=float(_cell(row, index, "Power (MW)")),
                    operating_years=_to_int(
                        _cell(row, index, "Operating life (years)"), default=20
                    ),
                    contract_structure=ContractStructure(
                        kind=contract_kind,
                        price_keur_per_mw_per_year=_to_float(
                            _cell(row, index, "Contract price (k€/MW/yr)"), default=0.0
                        ),
                        duration_years=_to_int(
                            _cell(row, index, "Contract duration (years)"), default=0
                        ),
                        revenue_sharing_above_floor_pct=_to_pct(
                            _cell(row, index, "Aggregator sharing above floor (%)"), default=0.0
                        ),
                    ),
                    gearing_pct_override=_to_pct(
                        _cell(row, index, "Gearing override (%)"), default=None
                    ),
                    interest_rate_override=_to_pct(
                        _cell(row, index, "Interest rate override (%)"), default=None
                    ),
                    dsa_keur_override=_to_float(
                        _cell(row, index, "DSA override (k€)"), default=None
                    ),
                    devex_keur_override=_to_float(
                        _cell(row, index, "DEVEX override (k€)"), default=None
                    ),
                    carry_months_override=_to_int(
                        _cell(row, index, "Carry months override"), default=None
                    ),
                    carry_rate_override=_to_pct(
                        _cell(row, index, "Carry rate override (%)"), default=None
                    ),
                    connection_capex_mode=connection_capex_mode,
                    manual_connection_capex_keur=_to_float(
                        _cell(row, index, "Manual grid connection CAPEX (k€)"), default=0.0
                    ),
                    distance_rte_km=_to_float(
                        _cell(row, index, "Distance to RTE substation (km)"), default=0.0
                    ),
                    land_lease_opex_keur=_to_float(
                        _cell(row, index, "Land lease OPEX (k€/yr)"), default=0.0
                    ),
                    land_lease_indexation_pct=_to_pct(
                        _cell(row, index, "Land lease indexation (%/yr)"), default=0.0
                    ),
                    oro_requested=oro_requested,
                    curtailment_hours=_to_int(
                        _cell(row, index, "ORO curtailment hours"), default=None
                    ),
                    repowering_enabled=repowering_enabled,
                    repowering_year_mode=repowering_year_mode,
                    repowering_op_year_manual=_to_int(
                        _cell(row, index, "Repowering year (manual)"), default=15
                    ),
                    turpe_50pct_reduction=_to_bool(
                        _cell(row, index, "TURPE 50% reduction"), default=False
                    ),
                )
            )
        except (ValueError, KeyError) as exc:
            row_label = f"'{name}'" if not _is_blank(name) else "(no name)"
            raise ValueError(f"Row {row_number} {row_label}: {exc}") from exc

    if not configs:
        raise ValueError(f"No project row found in the '{SHEET_NAME}' sheet.")
    return configs


def build_template_workbook() -> openpyxl.Workbook:
    """Classeur telechargeable depuis le Configurateur : onglet `Projects`
    (en-tetes + 2 lignes d'exemple) + onglet `Legend` (1 ligne par colonne,
    description/valeurs attendues) - pour que l'utilisateur puisse le
    pre-remplir sans deviner le format."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(
        [
            "Site A",
            2,
            "HTA",
            "Classique",
            "No",
            2028,
            10,
            20,
            "Full merchant",
            None,
            None,
            None,
            "No",
            None,
            "Yes",
            "manual",
            15,
            "library",
            None,
            None,
            0,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "No",
        ]
    )
    ws.append(
        [
            "Site B",
            4,
            "HTB2",
            "Injection",
            "Yes",
            2030,  # gabarit+ORO on 4h HTB2 is only modelled by Aurora for COD2030
            20,
            25,
            "Floor",
            80,
            10,
            0,
            "Yes",
            3000,
            "Yes",
            "auto",
            None,
            "distance_rte",
            None,
            5,
            250,
            2,
            70,
            5,
            None,
            None,
            None,
            None,
            "Yes",
        ]
    )
    for i in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    legend = wb.create_sheet(LEGEND_SHEET_NAME)
    legend.append(["Column", "When used", "Expected values"])
    for cell in legend[1]:
        cell.font = Font(bold=True)
    for label, when, expected in _COLUMN_SPECS:
        legend.append([label, when, expected])
    for i in range(1, 4):
        legend.column_dimensions[get_column_letter(i)].width = 36

    return wb


def template_bytes() -> bytes:
    buffer = io.BytesIO()
    build_template_workbook().save(buffer)
    return buffer.getvalue()
