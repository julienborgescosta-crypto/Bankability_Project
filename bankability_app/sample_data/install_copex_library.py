"""Installe une nouvelle version de la COPEX Library QEF dans `config/copex_icp.xlsx`.

Usage :
    .venv\\Scripts\\python.exe sample_data/install_copex_library.py "<COPEX LIBRARY.xlsx>"
        [--dso-grid-connection-keur 300]

- Retire les liens externes (references vers d'autres classeurs : chemins OneDrive/SharePoint de
  BP et quelques cellules en cache) - `config/copex_icp.xlsx` est committe, jamais un BP reel.
- Remplace dans les formules chaque reference externe par sa valeur en cache :
  `=794.13*'[2]I-Project'!$G$28^(-0.61)*1000` devient `=794.13*80^(-0.61)*1000`. L'app relit la
  loi de puissance dans le texte de la formule (`core/copex_icp.py`).
- Supprime les noms definis qui pointent vers un classeur externe.
- `--dso-grid-connection-keur` : remplace le raccordement DSO (la V1_20261001 indique 0,3 k€ au
  lieu de 300 k€).
- Relit le resultat avec `core.copex_icp` et affiche les valeurs cles.

Travaille directement sur le XML du classeur : un enregistrement par openpyxl effacerait les
valeurs en cache des formules.
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

APP_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = APP_ROOT / "config" / "copex_icp.xlsx"

_EXTERNAL_LINK_TYPE = "relationships/externalLink"
_EXTERNAL_REF = re.compile(
    r"(?:'\[(?P<qn>\d+)\](?P<qsheet>[^']+)'|\[(?P<n>\d+)\](?P<sheet>[A-Za-z0-9_.\-]+))"
    r"!\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)"
)


def _external_link_caches(parts: dict[str, bytes]) -> dict[int, dict[tuple[str, str], str]]:
    """{n de la reference [n]: {(onglet, cellule): valeur en cache}}, dans l'ordre des
    `<externalReference>` du classeur."""
    workbook = parts["xl/workbook.xml"].decode("utf-8")
    rels = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    targets = {
        m["id"]: m["target"]
        for m in re.finditer(
            r'<Relationship Id="(?P<id>[^"]+)"[^>]*Target="(?P<target>[^"]+)"', rels
        )
    }
    caches: dict[int, dict[tuple[str, str], str]] = {}
    for n, rel_id in enumerate(re.findall(r'<externalReference r:id="([^"]+)"', workbook), start=1):
        xml = parts["xl/" + targets[rel_id]].decode("utf-8")
        sheet_names = re.findall(r'<sheetName val="([^"]+)"', xml)
        cache: dict[tuple[str, str], str] = {}
        # Un onglet sans cellule en cache s'ecrit <sheetData sheetId="3"/>.
        for m in re.finditer(
            r'<sheetData sheetId="(?P<id>\d+)"[^>/]*(?:/>|>(?P<body>.*?)</sheetData>)', xml, re.S
        ):
            for ref, value in re.findall(
                r'<cell r="([A-Z]+\d+)"[^>]*>\s*<v>([^<]*)</v>', m["body"] or ""
            ):
                cache[(sheet_names[int(m["id"])], ref)] = value
        caches[n] = cache
    return caches


def _inline_external_refs(formula: str, caches: dict[int, dict[tuple[str, str], str]]) -> str:
    def replace(match: re.Match) -> str:
        n = int(match["qn"] or match["n"])
        sheet = match["qsheet"] or match["sheet"]
        key = (sheet, f"{match['col']}{match['row']}")
        if key not in caches.get(n, {}):
            raise SystemExit(
                f"Reference externe {match[0]} sans valeur en cache - impossible de la remplacer."
            )
        return caches[n][key]

    return _EXTERNAL_REF.sub(replace, formula)


def _dso_grid_connection_cells(source: Path) -> list[str]:
    ws = openpyxl.load_workbook(source, data_only=True, keep_links=False)["COPEX_library"]
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Case")
    row = next(
        r
        for r in range(header_row, ws.max_row + 1)
        if isinstance(ws.cell(r, 1).value, str) and "Grid connection" in ws.cell(r, 1).value
    )
    return [
        ws.cell(row, col).coordinate
        for col in range(2, ws.max_column + 1)
        if ws.cell(header_row, col).value == "DSO"
    ]


def install(source: Path, dso_grid_connection_keur: float | None) -> None:
    with zipfile.ZipFile(source) as zin:
        infos = zin.infolist()
        parts = {info.filename: zin.read(info.filename) for info in infos}
    caches = _external_link_caches(parts)

    def sheet_xml(xml: str) -> str:
        return re.sub(
            r"(<f\b[^>]*>)([^<]*)(</f>)",
            lambda m: m[1] + _inline_external_refs(m[2], caches) + m[3],
            xml,
        )

    replaced: dict[str, bytes] = {}
    for name, data in parts.items():
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            replaced[name] = sheet_xml(data.decode("utf-8")).encode("utf-8")

    workbook = parts["xl/workbook.xml"].decode("utf-8")
    workbook = re.sub(r"<externalReferences>.*?</externalReferences>", "", workbook, flags=re.S)
    workbook = re.sub(r"<definedName\b[^>]*>[^<]*\[\d+\][^<]*</definedName>", "", workbook)
    replaced["xl/workbook.xml"] = workbook.encode("utf-8")
    rels = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    replaced["xl/_rels/workbook.xml.rels"] = re.sub(
        rf'<Relationship [^>]*{_EXTERNAL_LINK_TYPE}"[^>]*/>', "", rels
    ).encode("utf-8")
    content_types = parts["[Content_Types].xml"].decode("utf-8")
    replaced["[Content_Types].xml"] = re.sub(
        r'<Override PartName="/xl/externalLinks/[^"]+"[^>]*/>', "", content_types
    ).encode("utf-8")

    if dso_grid_connection_keur is not None:
        sheet_name = "xl/worksheets/sheet1.xml"
        xml = replaced[sheet_name].decode("utf-8")
        for ref in _dso_grid_connection_cells(source):
            xml, count = re.subn(
                rf'(<c r="{ref}"[^>]*>(?:<f>[^<]*</f>)?<v>)[^<]*(</v>)',
                rf"\g<1>{dso_grid_connection_keur:g}\g<2>",
                xml,
            )
            if count != 1:
                raise SystemExit(f"Cellule {ref} (raccordement DSO) introuvable dans {sheet_name}.")
        replaced[sheet_name] = xml.encode("utf-8")

    with zipfile.ZipFile(OUTPUT_PATH, "w") as zout:
        for info in infos:
            if info.filename.startswith("xl/externalLinks/"):
                continue
            zout.writestr(info, replaced.get(info.filename, parts[info.filename]))
    print(f"Ecrit : {OUTPUT_PATH} ({len(caches)} lien(s) externe(s) retire(s))")


def report() -> None:
    sys.path.insert(0, str(APP_ROOT))
    from core import copex_icp

    library = copex_icp.load_icp_library(OUTPUT_PATH)
    print(f"Version : {library.version}")
    for tension, segment in copex_icp.TENSION_TO_ICP_SEGMENT.items():
        grid = library.line_items["Grid connection"][segment]
        print(
            f"  {tension} ({segment}) : batteries+PCS 2h "
            f"{library.line_items[copex_icp.battery_label(2)][segment].value:,.0f} €/MWh, "
            f"raccordement {grid.value:,.1f} ({grid.unit})"
        )
        if grid.unit == "keur_flat" and grid.value < 1.0:
            print(f"  ATTENTION : raccordement {segment} a {grid.value} k€ - erreur de saisie ?")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source", type=Path)
    parser.add_argument("--dso-grid-connection-keur", type=float, default=None)
    args = parser.parse_args()
    install(args.source, args.dso_grid_connection_keur)
    report()
